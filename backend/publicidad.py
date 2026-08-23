"""📣 MÓDULO PUBLICIDAD (solo Admin): listados de campaña, envío masivo de
templates por correo con protección de reputación, y campañas por WhatsApp."""
import re
import uuid
import asyncio
import logging
from pathlib import Path
from datetime import datetime, timezone
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Request, UploadFile, File, Form
from database import db

pub = APIRouter(prefix="/publicidad")
PUBLIC_DIR = Path("/app/frontend/public")
TEMPLATES = [
    {"archivo": "template-brokers-concreces.html", "nombre": "Brokers — dorado (oficial)"},
    {"archivo": "template-inmobiliarias-concreces.html", "nombre": "Inmobiliarias — dorado (bloques)"},
    {"archivo": "template-inmobiliarias-corporativo.html", "nombre": "Inmobiliarias — corporativo sobrio"},
    {"archivo": "template-clientes-directos-c.html", "nombre": "Clientes Directos — carta corporativa",
     "asunto": "Crédito Hipotecario / Viviendas Nuevas y Usadas con Entrega Inmediata"},
]
RX_MAIL = re.compile(r"^[\w.\-+]+@[\w\-]+(\.[\w\-]+)+$")
PAUSA_SEG = 6


def _now():
    return datetime.now(timezone.utc).isoformat()


def _exigir_admin(request):
    c = getattr(request.state, "user", {}) or {}
    if c.get("rol") not in ("admin", "maestro"):
        raise HTTPException(status_code=403, detail="Módulo Publicidad: exclusivo del Administrador")
    return c


def _parsear_contactos(texto, excluir):
    crudos = re.split(r"[\s,;<>\n\r\t]+", texto or "")
    vistos, contactos, invalidos, excluidos, dups = set(), [], [], [], 0
    for c in crudos:
        c = c.strip().strip('"\'').lower()
        if not c:
            continue
        tel = re.sub(r"[^\d+]", "", c)
        if RX_MAIL.match(c):
            valor, tipo = c, "correo"
        elif re.match(r"^\+?\d{9,15}$", tel):
            valor, tipo = tel, "telefono"
        else:
            invalidos.append(c)
            continue
        if valor in vistos:
            dups += 1
            continue
        vistos.add(valor)
        if any(ex and ex.lower() in valor for ex in (excluir or [])):
            excluidos.append(valor)
            continue
        contactos.append({"valor": valor, "tipo": tipo})
    return contactos, {"agregados": len(contactos), "duplicados_eliminados": dups,
                       "excluidos": excluidos, "invalidos": invalidos}


@pub.get("/listados")
async def listados(request: Request):
    _exigir_admin(request)
    regs = await db.publicidad_listados.find({}, {"_id": 0}).sort("creado", -1).to_list(100)
    return {"listados": regs, "templates": TEMPLATES}


@pub.post("/listados")
async def crear_listado(payload: dict, request: Request):
    u = _exigir_admin(request)
    nombre = (payload.get("nombre") or "").strip()
    if not nombre:
        raise HTTPException(status_code=400, detail="Nombre del listado obligatorio")
    contactos, resumen = _parsear_contactos(payload.get("contactos_texto") or "", payload.get("excluir") or [])
    existente = await db.publicidad_listados.find_one({"nombre": nombre}, {"_id": 0})
    if existente:
        previos = {c["valor"] for c in existente.get("contactos", [])}
        nuevos = [c for c in contactos if c["valor"] not in previos]
        resumen["duplicados_eliminados"] += len(contactos) - len(nuevos)
        resumen["agregados"] = len(nuevos)
        await db.publicidad_listados.update_one({"id": existente["id"]}, {
            "$push": {"contactos": {"$each": nuevos}}, "$set": {"actualizado": _now()}})
        lid = existente["id"]
    else:
        lid = str(uuid.uuid4())
        await db.publicidad_listados.insert_one({
            "id": lid, "nombre": nombre, "tipo_contacto": payload.get("tipo_contacto") or "Inmobiliaria / Empresa",
            "contactos": contactos, "creado": _now(), "creado_por": u.get("sub", "")})
    reg = await db.publicidad_listados.find_one({"id": lid}, {"_id": 0})
    return {"ok": True, "listado": reg, "resumen": resumen}


@pub.delete("/listados/{lid}")
async def borrar_listado(lid: str, request: Request):
    _exigir_admin(request)
    await db.publicidad_listados.delete_one({"id": lid})
    return {"ok": True}


TIPOS_DESTINATARIO = {"broker_inmobiliario": "Broker Inmobiliario",
                      "cliente_directo": "Cliente Directo",
                      "cliente_individual": "Cliente Individual"}


@pub.post("/listados/importar")
async def importar_listado(request: Request, archivo: UploadFile = File(...),
                           nombre: str = Form(""), excluir: str = Form("ecomac.cl"),
                           tipo_destinatario: str = Form("broker_inmobiliario")):
    """Carga un listado desde Excel (.xlsx) o CSV/TXT: extrae correos y teléfonos de todas
    las celdas (pueden venir en columnas distintas del MISMO archivo), deduplica y
    distribuye automáticamente: correos → campaña de mail · teléfonos → campaña WhatsApp."""
    u = _exigir_admin(request)
    if tipo_destinatario not in TIPOS_DESTINATARIO:
        tipo_destinatario = "broker_inmobiliario"
    raw = await archivo.read()
    fn = (archivo.filename or "").lower()
    celdas = []
    if fn.endswith((".xlsx", ".xlsm")):
        import io
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    celdas += [str(c) for c in row if c is not None]
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"No pude leer el Excel: {str(e)[:120]}")
    elif fn.endswith((".csv", ".txt")):
        celdas = [raw.decode("utf-8", errors="ignore")]
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado: suba .xlsx, .csv o .txt")
    exclusiones = [e.strip() for e in (excluir or "").split(",") if e.strip()]
    contactos, resumen = _parsear_contactos(" ".join(celdas), exclusiones)
    if not contactos:
        raise HTTPException(status_code=400,
                            detail="El archivo no contiene correos ni teléfonos válidos")
    nombre = (nombre or "").strip() or (archivo.filename or "Listado importado").rsplit(".", 1)[0]
    existente = await db.publicidad_listados.find_one({"nombre": nombre}, {"_id": 0})
    if existente:
        previos = {c["valor"] for c in existente.get("contactos", [])}
        nuevos = [c for c in contactos if c["valor"] not in previos]
        resumen["duplicados_eliminados"] += len(contactos) - len(nuevos)
        resumen["agregados"] = len(nuevos)
        await db.publicidad_listados.update_one({"id": existente["id"]}, {
            "$push": {"contactos": {"$each": nuevos}},
            "$set": {"actualizado": _now(), "tipo_destinatario": tipo_destinatario}})
        lid = existente["id"]
    else:
        lid = str(uuid.uuid4())
        await db.publicidad_listados.insert_one({
            "id": lid, "nombre": nombre, "tipo_contacto": "Importado (archivo)",
            "tipo_destinatario": tipo_destinatario,
            "contactos": contactos, "creado": _now(), "creado_por": u.get("sub", ""),
            "archivo_origen": archivo.filename})
    reg = await db.publicidad_listados.find_one({"id": lid}, {"_id": 0})
    n_mail = sum(1 for c in reg.get("contactos", []) if c["tipo"] == "correo")
    n_tel = sum(1 for c in reg.get("contactos", []) if c["tipo"] == "telefono")
    return {"ok": True, "listado": reg, "resumen": resumen,
            "mensaje": (f"«{nombre}» ({TIPOS_DESTINATARIO[tipo_destinatario]}): "
                        f"{resumen['agregados']} contacto(s) nuevos — distribución automática: "
                        f"{n_mail} correo(s) → Campañas de Correo · {n_tel} teléfono(s) → Campañas WhatsApp")}


# ══ CENTRO DE CAPTACIÓN — datos unificados para la vista del administrador ══
@pub.get("/captacion")
async def captacion(request: Request):
    _exigir_admin(request)
    prospectos = []
    async for p in db.prospectos.find({}, {"_id": 0}).sort("creado_en", -1).limit(60):
        oid = p.get("id")
        docs = await db.capturas_autonomas.count_documents({"oportunidad_id": oid})
        primer = (p.get("nombre") or "").split()[0].title() if p.get("nombre") else ""
        link = p.get("link_calificar") or ""
        mensaje = ""
        if link:
            mensaje = (f"🏠 *Central Mutuos - Precalificación Hipotecaria*\n\nHola {primer}, "
                       f"soy José Martín de Central Mutuos. Suba su Cédula y sus últimas 6 "
                       f"Liquidaciones de Sueldo en este portal privado y su calificación queda lista:\n{link}"
                       f"\n\nAtentamente, el equipo de @CentralMutuos")
        prospectos.append({"id": oid, "nombre": p.get("nombre") or "—",
                           "proyecto": p.get("proyecto") or "", "telefono": p.get("telefono") or "",
                           "estado": p.get("status") or "", "link": link,
                           "docs_subidos": docs, "captura_en": p.get("captura_autonoma_en") or "",
                           "mensaje_whatsapp": mensaje})
    llamadas = await db.solicitudes_llamada.find({}, {"_id": 0}).sort("creado_en", -1).limit(20).to_list(20)
    return {"prospectos": prospectos, "llamadas": llamadas,
            "capturas_total": await db.capturas_autonomas.count_documents({})}


@pub.get("/pendientes")
async def pendientes(request: Request):
    _exigir_admin(request)
    import whatsapp_twilio_service as wa
    ds19 = await db.publicidad_listados.find_one(
        {"nombre": {"$regex": "ds19", "$options": "i"}}, {"_id": 0, "nombre": 1, "contactos": 1})
    return {"ds19": {"resuelto": bool(ds19), "nombre": (ds19 or {}).get("nombre", ""),
                     "contactos": len((ds19 or {}).get("contactos", [])),
                     "detalle": "Listado 'ds19 01 inmobiliarias' (174 proyectos) pendiente: "
                                "las capturas de imagen eran ilegibles. Suba aquí el Excel/CSV de usatusubsidio.cl."},
            "twilio": {"resuelto": wa.configurado(),
                       "detalle": "WhatsApp automático (Regla de Oro #21): faltan las credenciales "
                                  "Twilio (Account SID, Auth Token y número exclusivo)."}}


@pub.post("/listados/{lid}/quitar")
async def quitar_contacto(lid: str, payload: dict, request: Request):
    _exigir_admin(request)
    await db.publicidad_listados.update_one({"id": lid}, {
        "$pull": {"contactos": {"valor": (payload.get("valor") or "").lower()}}})
    return {"ok": True}


async def _envio_bg(eid, correos, html, asunto, texto=""):
    import email_service as mail
    enviados, fallidos = 0, []
    for i, correo in enumerate(correos):
        try:
            r = await asyncio.to_thread(mail.send_mail, correo, asunto, html, None, "secundaria",
                                        body_text=texto or None)
            if r.get("success"):
                enviados += 1
            else:
                fallidos.append({"correo": correo, "error": str(r.get("error"))[:120]})
        except Exception as e:
            fallidos.append({"correo": correo, "error": str(e)[:120]})
        await db.publicidad_envios.update_one({"id": eid}, {"$set": {
            "enviados": enviados, "fallidos": fallidos, "progreso": i + 1}})
        if i < len(correos) - 1:
            await asyncio.sleep(PAUSA_SEG)
    await db.publicidad_envios.update_one({"id": eid}, {"$set": {
        "estado": "terminado", "terminado": _now()}})
    logging.info(f"📣 Campaña {eid}: {enviados}/{len(correos)} enviados, {len(fallidos)} fallidos")


def _app_url():
    try:
        for line in open("/app/frontend/.env"):
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().strip('"').rstrip("/")
    except Exception:
        pass
    return ""


def _render_campana(html):
    """Reemplaza los placeholders del template: personalización + link del formulario público."""
    return (html.replace("{{NOMBRE}}", "cliente")
                .replace("{{LINK_CONTACTO}}", f"{_app_url()}/api/publicidad/contacto"))


def _texto_plano(html):
    """Versión texto plano alternativa (anti-spam): mismo contenido sin HTML."""
    t = re.sub(r"<style.*?</style>", " ", html, flags=re.S | re.I)
    t = re.sub(r"<br\s*/?>", "\n", t, flags=re.I)
    t = re.sub(r"</(p|tr|table|div|h1|h2)>", "\n", t, flags=re.I)
    t = re.sub(r"<a[^>]*href=[\"']([^\"']+)[\"'][^>]*>([^<]*)</a>", r"\2: \1", t, flags=re.I)
    t = re.sub(r"<[^>]+>", " ", t)
    t = t.replace("&nbsp;", " ").replace("&amp;", "&")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n\s+", "\n", t)
    return t.strip()


@pub.post("/enviar")
async def enviar(payload: dict, request: Request):
    u = _exigir_admin(request)
    template = payload.get("template") or ""
    if template not in [t["archivo"] for t in TEMPLATES]:
        raise HTTPException(status_code=400, detail="Template inválido")
    asunto = (payload.get("asunto") or "").strip()
    if not asunto:
        raise HTTPException(status_code=400, detail="Asunto obligatorio")
    html = _render_campana((PUBLIC_DIR / template).read_text())
    texto = _texto_plano(html)
    if payload.get("prueba"):
        import email_service as mail
        r = await asyncio.to_thread(mail.send_mail, "ethangerardobarr@gmail.com",
                                    f"[PRUEBA] {asunto}", html, None, "secundaria",
                                    body_text=texto)
        if not r.get("success"):
            raise HTTPException(status_code=500, detail=f"Prueba falló: {str(r.get('error'))[:120]}")
        return {"ok": True, "prueba": True,
                "mensaje": "Correo de PRUEBA enviado a tu bandeja (gerardo.ext). Revísalo antes de la campaña real."}
    listado = await db.publicidad_listados.find_one({"id": payload.get("listado_id")}, {"_id": 0})
    if not listado:
        raise HTTPException(status_code=404, detail="Listado no encontrado")
    correos = [c["valor"] for c in listado.get("contactos", []) if c["tipo"] == "correo"]
    if not correos:
        raise HTTPException(status_code=400, detail="El listado no tiene correos válidos")
    if not payload.get("confirmado"):
        raise HTTPException(status_code=400, detail=f"Confirma el envío: se despachará a {len(correos)} destinatario(s)")
    eid = str(uuid.uuid4())
    await db.publicidad_envios.insert_one({
        "id": eid, "canal": "correo", "listado": listado["nombre"], "listado_id": listado["id"],
        "template": template, "asunto": asunto, "total": len(correos), "enviados": 0,
        "fallidos": [], "progreso": 0, "estado": "enviando", "iniciado": _now(),
        "por": u.get("sub", "")})
    asyncio.create_task(_envio_bg(eid, correos, html, asunto, texto))
    return {"ok": True, "envio_id": eid, "total": len(correos),
            "mensaje": f"Campaña iniciada: {len(correos)} correo(s) en segundo plano con pausa de {PAUSA_SEG}s entre envíos (protección de reputación)."}


@pub.get("/envios")
async def envios(request: Request):
    _exigir_admin(request)
    regs = await db.publicidad_envios.find({}, {"_id": 0}).sort("iniciado", -1).to_list(40)
    return {"envios": regs}


@pub.post("/whatsapp-links")
async def whatsapp_links(payload: dict, request: Request):
    u = _exigir_admin(request)
    mensaje = (payload.get("mensaje") or "").strip()
    if not mensaje:
        raise HTTPException(status_code=400, detail="Mensaje obligatorio")
    listado = await db.publicidad_listados.find_one({"id": payload.get("listado_id")}, {"_id": 0})
    if not listado:
        raise HTTPException(status_code=404, detail="Listado no encontrado")
    tels = [c["valor"] for c in listado.get("contactos", []) if c["tipo"] == "telefono"]
    if not tels:
        raise HTTPException(status_code=400, detail="El listado no tiene teléfonos: agrega números (+569XXXXXXXX) al listado")
    links = [{"telefono": t, "link": f"https://wa.me/{t.lstrip('+')}?text={quote(mensaje)}"} for t in tels]
    await db.publicidad_envios.insert_one({
        "id": str(uuid.uuid4()), "canal": "whatsapp", "listado": listado["nombre"],
        "listado_id": listado["id"], "asunto": mensaje[:80], "total": len(tels),
        "enviados": 0, "fallidos": [], "estado": "manual", "iniciado": _now(), "por": u.get("sub", "")})
    return {"ok": True, "links": links}


# ══ FORMULARIO PÚBLICO "QUIERO SER CONTACTADO" (campañas clientes directos) ══
FORM_CONTACTO_HTML = """<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Central Mutuos — Quiero ser contactado</title></head>
<body style="margin:0;background:#ededee;font-family:Georgia,'Times New Roman',serif;">
<div style="max-width:480px;margin:40px auto;background:#fff;">
  <div style="background:#0e0e10;padding:26px 34px;">
    <div style="font-size:20px;letter-spacing:5px;color:#d4af37;">CENTRAL MUTUOS</div>
    <div style="font-size:10px;letter-spacing:4px;color:#f5e7b8;margin-top:4px;">CON CRECES</div>
  </div>
  <div style="height:4px;background:#d4af37;"></div>
  <div style="padding:30px 34px 34px;">
    <h1 style="margin:0 0 6px;font-size:19px;color:#111;font-weight:normal;">Quiero ser contactado</h1>
    <div style="height:2px;background:#d4af37;width:60px;margin:0 0 16px;"></div>
    <p style="margin:0 0 20px;font-size:13.5px;color:#3a3a3a;line-height:1.6;">Déjanos tu nombre y teléfono y un ejecutivo de Central Mutuos te contactará para orientarte en tu crédito hipotecario.</p>
    <form id="f">
      <label style="display:block;font-size:12px;color:#8a6a0f;letter-spacing:1px;margin-bottom:4px;">NOMBRE</label>
      <input id="nombre" required minlength="3" maxlength="80" style="width:100%;box-sizing:border-box;padding:11px 12px;border:1px solid #cfc7ad;font-family:Georgia,serif;font-size:14px;margin-bottom:14px;" placeholder="Tu nombre completo">
      <label style="display:block;font-size:12px;color:#8a6a0f;letter-spacing:1px;margin-bottom:4px;">TEL&Eacute;FONO</label>
      <input id="telefono" required style="width:100%;box-sizing:border-box;padding:11px 12px;border:1px solid #cfc7ad;font-family:Georgia,serif;font-size:14px;margin-bottom:20px;" placeholder="+56 9 XXXX XXXX">
      <button id="btn" type="submit" style="width:100%;background:#0e0e10;color:#d4af37;border:none;border-bottom:2px solid #d4af37;padding:13px;font-family:Georgia,serif;font-size:14px;letter-spacing:1px;cursor:pointer;">Enviar solicitud</button>
    </form>
    <div id="msg" style="display:none;margin-top:16px;padding:12px 14px;font-size:13px;line-height:1.5;"></div>
    <p style="margin:22px 0 0;font-size:10px;color:#999;line-height:1.5;">Aprobación sujeta a revisión de antecedentes crediticios según la normativa vigente y las políticas de crédito internas de nuestra compañía. Mutuaria regulada por la CMF.</p>
  </div>
</div>
<script>
document.getElementById('f').addEventListener('submit',async function(e){
  e.preventDefault();
  var b=document.getElementById('btn'),m=document.getElementById('msg');
  b.disabled=true;b.textContent='Enviando…';
  try{
    var r=await fetch(window.location.pathname,{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({nombre:document.getElementById('nombre').value,telefono:document.getElementById('telefono').value})});
    var d=await r.json();m.style.display='block';
    if(r.ok&&d.ok){m.style.background='#f0f7f1';m.style.border='1px solid #2e7d43';m.style.color='#2e7d43';
      m.textContent=d.mensaje;document.getElementById('f').style.display='none';}
    else{m.style.background='#fdf1f1';m.style.border='1px solid #b91c1c';m.style.color='#b91c1c';
      m.textContent=(d.detail||'No se pudo enviar. Intenta nuevamente.');b.disabled=false;b.textContent='Enviar solicitud';}
  }catch(err){m.style.display='block';m.style.background='#fdf1f1';m.style.border='1px solid #b91c1c';
    m.style.color='#b91c1c';m.textContent='Error de conexión. Intenta nuevamente.';b.disabled=false;b.textContent='Enviar solicitud';}
});
</script>
</body></html>"""


@pub.get("/contacto")
async def contacto_form():
    from fastapi.responses import HTMLResponse
    return HTMLResponse(FORM_CONTACTO_HTML)


@pub.post("/contacto")
async def contacto_submit(payload: dict):
    nombre = (payload.get("nombre") or "").strip()[:80]
    tel = re.sub(r"[^\d+]", "", payload.get("telefono") or "")
    if len(nombre) < 3:
        raise HTTPException(status_code=400, detail="Ingresa tu nombre completo")
    if not re.match(r"^\+?\d{8,15}$", tel):
        raise HTTPException(status_code=400, detail="Ingresa un teléfono válido (+56 9 XXXX XXXX)")
    ya = await db.solicitudes_llamada.find_one({
        "telefono": tel, "origen": "campania_clientes_directos",
        "creado_en": {"$gt": (datetime.now(timezone.utc) - __import__("datetime").timedelta(hours=24)).isoformat()}})
    if not ya:
        await db.solicitudes_llamada.insert_one({
            "id": str(uuid.uuid4()), "cliente": nombre, "telefono": tel,
            "horario": "Cuanto antes", "motivo": "campaña clientes directos",
            "origen": "campania_clientes_directos", "creado_en": _now()})
        logging.info(f"📣 Contacto de campaña recibido: {nombre} · {tel}")
    return {"ok": True, "mensaje": ("¡Gracias! Recibimos tu solicitud. Un ejecutivo de Central Mutuos "
                                    "te contactará a la brevedad.")}
